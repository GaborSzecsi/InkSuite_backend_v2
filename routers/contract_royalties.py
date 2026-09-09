"""Render first-right royalties without a fixed number of template slots."""
from copy import deepcopy
from decimal import Decimal, InvalidOperation
import re
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph


class RoyaltyValidationError(ValueError):
    pass


def _number(value, label, *, percent=False):
    if value is None or isinstance(value, bool) or str(value).strip() == "":
        raise RoyaltyValidationError(f"{label} is missing. Complete the royalty builder before generating.")
    try:
        n = Decimal(str(value))
    except InvalidOperation:
        raise RoyaltyValidationError(f"{label} must be a number.")
    if not n.is_finite() or n < 0 or (percent and n > 100):
        raise RoyaltyValidationError(f"{label} is outside its allowed range.")
    return format(n, 'f').rstrip('0').rstrip('.') if '.' in format(n, 'f') else str(n)


def _key(value):
    return re.sub(r'[\s_-]', '', str(value)).lower()


def _roman(n):
    result = ''
    for value, text in [(1000,'m'),(900,'cm'),(500,'d'),(400,'cd'),(100,'c'),(90,'xc'),(50,'l'),(40,'xl'),(10,'x'),(9,'ix'),(5,'v'),(4,'iv'),(1,'i')]:
        while n >= value:
            result += text
            n -= value
    return result


def _condition(c, label):
    kind, op, raw = c.get('kind'), c.get('comparator'), c.get('value')
    if kind not in ('units', 'discount') or op not in ('<','<=','>','>=','between'):
        raise RoyaltyValidationError(f"{label}: unsupported royalty condition.")
    if op == 'between':
        if not isinstance(raw, (list, tuple)) or len(raw) != 2:
            raise RoyaltyValidationError(f"{label}: enter both range limits.")
        low, high = [_number(v, label, percent=kind == 'discount') for v in raw]
        if Decimal(low) >= Decimal(high):
            raise RoyaltyValidationError(f"{label}: range limits must increase.")
        if kind == 'discount':
            return f"sold at a discount of at least {low}% but less than {high}%"
        return f"for copies numbered {low} through {high}"
    n = _number(raw, label, percent=kind == 'discount')
    if kind == 'discount':
        wording = {'<':'less than','<=':'at or below','>':'greater than','>=':'at or above'}[op]
        return f"sold at a discount {wording} {n}%"
    if Decimal(n) != Decimal(n).to_integral_value():
        raise RoyaltyValidationError(f"{label}: copy limits must be whole numbers.")
    return {'<':f"for copies numbered less than {n}", '<=':f"on the first {n} copies sold", '>':f"on all copies sold over {n}", '>=':f"for copies numbered {n} and above"}[op]


def _discount_cap(tier, row):
    """Only the builder's unconditional, open-ended zero-rate discount tier moves."""
    rate = tier.get('rate_percent', tier.get('percent'))
    if rate is None or str(rate).strip() == '' or str(tier.get('note') or '').strip():
        return None
    try:
        if Decimal(str(rate)) != 0:
            return None
    except InvalidOperation:
        return None
    conditions = [c for c in (tier.get('conditions') or [])
                  if not (row.get('condition_mode') == 'all_copies' and c.get('kind') == 'units')]
    if len(conditions) != 1:
        return None
    c = conditions[0]
    if c.get('kind') != 'discount' or c.get('comparator') not in ('>', '>='):
        return None
    return (c['comparator'], _number(c.get('value'), f"{row.get('format')} no-royalty cutoff", percent=True))


def _no_royalty_clause(memo, party):
    rows = ((memo.get('royalties') or {}).get(party) or {}).get('first_rights') or []
    groups = {}
    count = 0
    for row in rows:
        tiers = row.get('tiers') or []
        if not (row.get('escalating') or row.get('mode') == 'tiered' or len(tiers) > 1):
            continue
        caps = {_discount_cap(t, row) for t in tiers} - {None}
        if len(caps) > 1:
            raise RoyaltyValidationError(f"{row.get('format')}: enter one no-royalty discount cutoff.")
        if caps:
            cap = caps.pop()
            groups.setdefault(cap, []).append(str(row.get('format')).lower())
            count += 1
    if not groups:
        return None
    def join(names):
        return names[0] if len(names) == 1 else ', '.join(names[:-1]) + ' and ' + names[-1]
    pieces = []
    for (op, cutoff), names in groups.items():
        formats = 'any format' if len(groups) == 1 and count == len(rows) else join(names) + (' edition' if len(names) == 1 else ' editions')
        comparison = 'at or above' if op == '>=' else 'greater than'
        pieces.append(f"copies of {formats} sold at a discount {comparison} {cutoff}%")
    return 'No royalties shall be payable on ' + '; or on '.join(pieces) + '.'


def build_royalty_clauses(memo, party='author', *, separate_discount_caps=False):
    rows = ((memo.get('royalties') or {}).get(party) or {}).get('first_rights') or []
    result = []
    seen = set()
    for row in rows:
        name = str(row.get('format') or '').strip()
        if not name:
            raise RoyaltyValidationError('A royalty format is missing.')
        if _key(name) in seen:
            raise RoyaltyValidationError(f'{name}: combine duplicate format rows before generating.')
        seen.add(_key(name))
        tiers = row.get('tiers') or []
        tiered = row.get('escalating') or row.get('mode') == 'tiered' or len(tiers) > 1
        if tiered and not tiers:
            raise RoyaltyValidationError(f'{name}: add at least one royalty tier.')
        terms = tiers if tiered else [{'rate_percent':row.get('flat_rate_percent',row.get('percent')), 'conditions':[], 'base':row.get('base')}]
        pieces = []
        for i, tier in enumerate(terms):
            label = f'{name}, tier {i+1}'
            rate = _number(tier.get('rate_percent',tier.get('percent')), f'{label} royalty percentage', percent=True)
            base = tier.get('base') or row.get('base')
            if base not in ('list_price','net_receipts'):
                raise RoyaltyValidationError(f'{label}: select List Price or Net Receipts.')
            basis = "the Book’s suggested retail price" if base == 'list_price' else 'the net amount received by the Publisher'
            conditions = tier.get('conditions') or []
            # The builder ignores copy thresholds in all-copies mode.
            conditions = [c for c in conditions if not (row.get('condition_mode') == 'all_copies' and c.get('kind') == 'units')]
            units = [c for c in conditions if c.get('kind') == 'units']
            other = [c for c in conditions if c.get('kind') != 'units']
            wording = ' and '.join(_condition(c, label) for c in units) if units else 'on all copies sold'
            if other:
                wording += ' ' + ' and '.join(_condition(c, label).removeprefix('sold ') for c in other)
            piece = f'{rate}% of {basis} {wording}'
            note = str(tier.get('note') or '').strip()
            if note:
                piece += f' ({note})'
            if not (separate_discount_caps and _discount_cap(tier, row)):
                pieces.append(piece)
        pieces = [(f'({_roman(i+1)}) ' if len(pieces) > 1 else '') + piece for i, piece in enumerate(pieces)]
        if not pieces:
            continue
        result.append((_key(name), f'On sales of the {name.lower()} edition of the Book: ' + '; '.join(pieces) + '.'))
    return result


_LEGACY = re.compile(r'\{\{\s*(Hardcover_[^}]+|Paperback_[^}]+|Boardbook_[^}]+|Ebook)\s*\}\}', re.I)
_BLOCK = re.compile(r'^\s*\{\{\s*ROYALTIES_BLOCK\s*\}\}\s*$', re.I)


def render_royalty_sections(doc, memo, party='author', insert_block=None):
    """Keep paragraph properties; preserve prose following the ebook rate sentence."""
    cutoff_clause = _no_royalty_clause(memo, party)
    clauses = build_royalty_clauses(memo, party, separate_discount_caps=bool(cutoff_clause))
    by_format = dict(clauses)
    containers = [doc]
    def tables(container):
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield cell
                    yield from tables(cell)
    containers.extend(tables(doc))
    if cutoff_clause:
        # Match the dedicated exemption paragraph, never unrelated discount prose.
        cutoff_pattern = re.compile(r'^On copies of any format sold at or higher than \d+(?:\.\d+)?% discounts\.$', re.I)
        targets = {}
        for container in containers:
            for p in container.paragraphs:
                if cutoff_pattern.fullmatch(p.text.strip()) or p.text.strip() == '{{NO_ROYALTY_DISCOUNT_BLOCK}}':
                    targets[p._p] = p
        if not targets:
            raise RoyaltyValidationError('Template needs a NO_ROYALTY_DISCOUNT_BLOCK placeholder in its no-royalty section to include the discount cutoff.')
        for p in targets.values():
            props = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
            p.clear()
            run = p.add_run(cutoff_clause)
            if props is not None:
                run._r.insert(0, props)
            run.font.color.rgb = RGBColor(255, 0, 0)
    seen_containers = set()
    for container in containers:
        element = getattr(container, '_tc', doc._element)
        if element in seen_containers:
            continue
        seen_containers.add(element)
        paragraphs = list(container.paragraphs)
        blocks = [p for p in paragraphs if _BLOCK.fullmatch(p.text)]
        legacy = [p for p in paragraphs if _LEGACY.search(p.text)]
        if blocks or legacy:
            # Each generated clause declares its own royalty basis (list price or net receipts).
            for heading in paragraphs:
                old = "the following percentages of the Book's suggested retail price:"
                if heading.text.strip().startswith('Royalties.') and old in heading.text:
                    # Replace within runs when possible, preserving the section numbering and style.
                    for run in heading.runs:
                        if old in run.text:
                            run.text = run.text.replace(old, 'the following royalties:')
                            break
                    else:
                        text = heading.text.replace(old, 'the following royalties:')
                        heading.clear()
                        heading.add_run(text)
        if blocks and legacy:
            raise RoyaltyValidationError('Template contains both ROYALTIES_BLOCK and legacy royalty slots. Use one royalty section.')
        if blocks:
            for p in blocks:
                if not insert_block or not insert_block(p, '\n'.join(text for _,text in clauses)):
                    raise RoyaltyValidationError('Place ROYALTIES_BLOCK below a numbered contract section heading.')
            continue
        if not legacy:
            continue
        applied = set()
        last = legacy[-1]
        def write(p, text, tail=''):
            props = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
            p.clear()
            run = p.add_run(text)
            if props is not None:
                run._r.insert(0, props)
            run.font.color.rgb = RGBColor(255,0,0)
            if tail:
                p.add_run(tail)
        for p in legacy:
            tokens = _LEGACY.findall(p.text)
            keys = {_key(t.split('_')[0]) for t in tokens}
            if len(keys) != 1:
                raise RoyaltyValidationError('Each royalty format must occupy its own template paragraph.')
            key = keys.pop()
            if key not in by_format:
                continue
            tail = ''
            if key == 'ebook':
                # Preserve the publisher's existing renegotiation/sharing provisions verbatim.
                end = p.text.find('.', _LEGACY.search(p.text).end())
                if end < 0:
                    raise RoyaltyValidationError('The ebook royalty sentence must end with a period before additional provisions.')
                tail = p.text[end+1:]
            write(p, by_format[key], tail)
            applied.add(key)
        # Additional formats inherit the existing royalty list's paragraph style/numbering.
        anchor = last._p
        for key, text in clauses:
            if key in applied:
                continue
            new = deepcopy(last._p)
            anchor.addnext(new)
            write(Paragraph(new, last._parent), text)
            anchor = new
        for p in legacy:
            match = _LEGACY.search(p.text)
            if match:  # Unselected format; matched paragraphs were already rewritten.
                p._p.getparent().remove(p._p)
