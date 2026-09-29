"""Run with python -m unittest discover -s tests -p test_illustrator_memo_regressions.py."""
import io
import re
import unittest
from copy import deepcopy
from contextlib import ExitStack
from unittest.mock import MagicMock, patch
from docx import Document
from docx.shared import Pt
from routers import deal_memo_drafts as drafts
from routers.contract_royalties import render_royalty_sections, RoyaltyValidationError

PREFIX = ("16. The Publisher shall pay the Illustrator no royalties on a commercially reasonable "
          "number of copies of the Work sold to the Illustrator, distributed for review, advertising, "
          "publicity, or sales promotion, sold at or below the cost of manufacture, or that have been "
          "damaged or destroyed. ")
SENTENCE = "No royalties on copies sold at a discount at, or greater than 75%;"
ROW = {"format": "Hardcover", "mode": "tiered", "base": "list_price", "tiers": [
    {"rate_percent": 10, "conditions": [{"kind": "discount", "comparator": "<", "value": 70}]},
    {"rate_percent": 0, "conditions": [{"kind": "discount", "comparator": ">=", "value": 70}]}]}

class IllustratorRegressions(unittest.TestCase):
    def test_inline_cutoff_preserves_exemptions_and_run_formatting(self):
        for role in ("author", "illustrator"):
            for in_table in (False, True):
                doc = Document()
                p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0] if in_table else doc.add_paragraph()
                original = p.add_run(PREFIX)
                original.font.name = "Garamond"
                original.font.size = Pt(11)
                original.italic = True
                original_xml = original._r.xml
                p.add_run(SENTENCE[:27]).bold = True
                p.add_run(SENTENCE[27:])
                tail = p.add_run(" Other provisions remain.")
                tail.underline = True
                tail_xml = tail._r.xml
                render_royalty_sections(doc, {"royalties": {role: {"first_rights": [ROW]}}}, role)
                self.assertEqual(p.text, PREFIX + "No royalties shall be payable on copies of any format sold at a discount at or above 70%; Other provisions remain.")
                self.assertEqual(original._r.xml, original_xml)
                self.assertEqual(tail._r.xml, tail_xml)
                saved = io.BytesIO()
                doc.save(saved)
                saved.seek(0)
                loaded = Document(saved)
                output = loaded.tables[0].cell(0, 0).paragraphs[0] if in_table else loaded.paragraphs[0]
                self.assertEqual(output.text, p.text)

    def test_legacy_and_explicit_block_still_work(self):
        for text in ("On copies of any format sold at or higher than 75% discounts.", "{{NO_ROYALTY_DISCOUNT_BLOCK}}"):
            doc = Document()
            p = doc.add_paragraph(text)
            render_royalty_sections(doc, {"royalties": {"illustrator": {"first_rights": [ROW]}}}, "illustrator")
            self.assertIn("at or above 70%", p.text)

    def test_unrelated_discount_prose_is_not_replaced(self):
        doc = Document()
        p = doc.add_paragraph("The Publisher may sell copies at a discount of 75%.")
        with self.assertRaises(RoyaltyValidationError):
            render_royalty_sections(doc, {"royalties": {"illustrator": {"first_rights": [ROW]}}}, "illustrator")
        self.assertEqual(p.text, "The Publisher may sell copies at a discount of 75%.")

    def test_agency_objects_are_never_names(self):
        self.assertEqual(drafts._person_name({"agency": {"name": "Agency"}}), "")
        self.assertEqual(drafts._person_name("{'agency': {'name': 'Agency'}}"), "")
        self.assertEqual(drafts._person_name({"display_name": "Illustrator Name", "agency": {}}), "Illustrator Name")

    def test_actual_royalty_and_advance_serializers_round_trip(self):
        class Cursor:
            def __init__(self):
                self.tables = {}
                self.result = []
            def execute(self, sql, params):
                insert = re.search(r"INSERT INTO (\w+)\s*\((.*?)\)\s*VALUES", sql, re.S)
                if insert:
                    table, columns = insert.groups()
                    names = [name.strip() for name in columns.split(",")]
                    record = dict(zip(names, params))
                    records = self.tables.setdefault(table, [])
                    record["id"] = f"{table}-{len(records)}"
                    records.append(record)
                    self.result = [record]
                else:
                    table = re.search(r"FROM (\w+)", sql).group(1)
                    self.result = self.tables.get(table, [])
            def fetchone(self):
                return self.result[0]
            def fetchall(self):
                return self.result
        cur = Cursor()
        body = {"royalties": {"illustrator": {"first_rights": [ROW], "subrights": []},
                              "author": {"first_rights": [{"format": "Ebook", "mode": "flat", "flat_rate_percent": 25, "base": "net_receipts"}], "subrights": []}},
                "advanceSchedule": [{"amountType": "amount", "value": 1250, "trigger": "Signing"}]}
        drafts._save_royalties(cur, "tenant", "draft", body)
        royalties = drafts._hydrate_royalties(cur, "draft")
        self.assertEqual(royalties["illustrator"]["first_rights"][0]["tiers"][1]["rate_percent"], 0)
        self.assertEqual(royalties["illustrator"]["first_rights"][0]["tiers"][1]["conditions"], ROW["tiers"][1]["conditions"])
        self.assertEqual(royalties["author"]["first_rights"][0]["flat_rate_percent"], 25)
        drafts._save_advance_schedule(cur, "tenant", "draft", body)
        schedule = drafts._hydrate_advance_schedule(cur, "draft")
        self.assertEqual(schedule[0]["value"], 1250)
        self.assertEqual(schedule[0]["amountType"], "amount")
        self.assertEqual(schedule[0]["trigger"], "Signing")

    def test_save_and_reopen_illustrator_fields(self):
        contributor = {"role_code": "A12", "role_label": "Illustrator", "display_name": "Illustrator Name", "name": "Illustrator Name", "email": "person@example.test", "address": {"street": "1 Test St"}}
        body = {"uid": "memo-test", "title": "Same Book", "contributorRole": "illustrator",
                "contributor_role": "author", "selectedTemplate": "illustrator-template",
                "selected_work_id": "work-test", "illustrator": contributor,
                "author": {"agency": {"name": "Agency"}}, "illustrator_advance": 2500,
                "advanceSchedule": [{"amountType": "percent", "value": 100, "trigger": "Signing"}],
                "royalties": {"author": {"first_rights": [], "subrights": []},
                              "illustrator": {"first_rights": [ROW], "subrights": []}}}
        stored = {}
        def execute(sql, params):
            if "INSERT INTO deal_memo_drafts" in sql:
                columns = sql.split("INSERT INTO deal_memo_drafts (", 1)[1].split(")", 1)[0]
                names = [c.strip() for c in columns.split(",")]
                self.assertEqual(len(names), len(params))
                stored.update(zip(names, params))
                stored["id"] = "draft-test"
        cur = MagicMock()
        cur.__enter__.return_value = cur
        cur.execute.side_effect = execute
        cur.fetchone.return_value = {"id": "draft-test"}
        conn = MagicMock()
        conn.__enter__.return_value = conn
        conn.cursor.return_value = cur
        with ExitStack() as stack:
            stack.enter_context(patch.object(drafts, "db_conn", return_value=conn))
            stack.enter_context(patch.object(drafts, "_require_tables"))
            stack.enter_context(patch.object(drafts, "_tenant_id", return_value="tenant-test"))
            stack.enter_context(patch.object(drafts, "_fetch_one_draft", side_effect=lambda *args: deepcopy(stored)))
            stack.enter_context(patch.object(drafts, "_clear_children"))
            for name in ("_save_advance_schedule", "_save_royalties", "_save_deal_memo_contributors"):
                stack.enter_context(patch.object(drafts, name))
            stack.enter_context(patch.object(drafts, "_hydrate_advance_schedule", return_value=body["advanceSchedule"]))
            stack.enter_context(patch.object(drafts, "_hydrate_royalties", return_value=body["royalties"]))
            stack.enter_context(patch.object(drafts, "_hydrate_deal_memo_contributors", return_value=[contributor]))
            drafts.upsert_deal_memo(body, "tenant-test")
            reopened = drafts.get_deal_memo("memo-test", "tenant-test")
        self.assertEqual(reopened["contributorRole"], "illustrator")
        self.assertEqual(reopened["contributor_role_code"], "A12")
        self.assertEqual(reopened["selectedTemplate"], "illustrator-template")
        self.assertEqual(reopened["selected_work_id"], "work-test")
        self.assertEqual(reopened["illustrator"]["display_name"], "Illustrator Name")
        self.assertEqual(reopened["illustrator_advance"], 2500)
        self.assertEqual(reopened["advanceSchedule"], body["advanceSchedule"])
        self.assertEqual(reopened["royalties"]["illustrator"], body["royalties"]["illustrator"])
        self.assertEqual(reopened["author"], "")

if __name__ == "__main__":
    unittest.main()
