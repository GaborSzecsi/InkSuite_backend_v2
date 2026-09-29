"""Contract identity tags, tested without importing storage/network initialization."""
import ast
import io
import re
import unittest
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor, Pt
from routers.contract_royalties import format_contract_insertion

SOURCE = Path(__file__).resolve().parents[1] / 'routers' / 'contract_docs.py'
tree = ast.parse(SOURCE.read_text(encoding='utf-8'))
ns = {'format_contract_insertion': format_contract_insertion, 're': re, 'RGBColor': RGBColor, 'TOKEN_RE': re.compile(r'\{\{\s*([^{}]+?)\s*\}\}')}
for name in ('_dig', '_first_non_empty', '_normalize_token_name', '_get_value_from_memo',
             '_default_mapping', '_replace_inline_tokens', '_iter_contract_paragraphs',
             '_build_art_delivery_block', '_build_manuscript_delivery_block', '_insert_numbered_contract_block'):
    node = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == name)
    exec(compile(ast.Module(body=[node], type_ignores=[]), str(SOURCE), 'exec'), ns)

class IdentityFields(unittest.TestCase):
    def test_palette_tags_and_legacy_author_tags(self):
        memo = {'illustrator_name': 'Iris Artist', 'illustrator_email': 'iris@example.test',
                'illustrator_phone_number': '555-0100', 'illustrator_address': {
                    'street': '1 Art Lane', 'city': 'Boston', 'state': 'MA', 'zip': '02101'},
                'agent_name': 'Selected Agent', 'agent_email': 'agent@example.test',
                'agency_name': 'Art Agency', 'agency_street': '2 Agency Road',
                'agency_phone_number': '555-0200', 'agency_email': 'agency@example.test'}
        expected = {'Illustrator_Name': 'Iris Artist', 'Illustrator_Email': 'iris@example.test',
                    'Illustrator_Phone': '555-0100', 'Illustrator_Street_Address': '1 Art Lane',
                    'Illustrator_City': 'Boston', 'Illustrator_State': 'MA', 'Illustrator_Zip': '02101',
                    'Agent_Name': 'Selected Agent', 'Agent_Email': 'agent@example.test',
                    'Agency_Name': 'Art Agency', 'Agency_Address': '2 Agency Road',
                    'Agency_Phone': '555-0200', 'Agency_Email': 'agency@example.test',
                    'Author_Name': 'Iris Artist', 'Author_Email': 'iris@example.test'}
        mapping = ns['_default_mapping']('illustrator')
        values = {token: ns['_get_value_from_memo'](memo, field) for token, field in mapping.items()}
        doc = Document()
        for token, value in expected.items():
            p = doc.add_paragraph('Before ')
            p.add_run('{{' + token[:5]).font.name = 'Garamond'
            p.add_run(token[5:] + '}}').font.name = 'Garamond'
            p.add_run(' after.').italic = True
            ns['_replace_inline_tokens'](p, values)
            self.assertEqual(p.text, 'Before ' + value + ' after.')
            self.assertTrue(p.runs[-1].italic)
            inserted = next(r for r in p.runs if r.text == value)
            self.assertEqual(inserted.font.name, 'Times New Roman')
            self.assertEqual(inserted.font.size, Pt(12))
        self.assertEqual(ns['_default_mapping']('author')['Author_Name'], 'author')
        self.assertEqual(ns['_get_value_from_memo']({'illustrator': {'display_name': 'Nested Artist'}}, 'illustrator_name'), 'Nested Artist')

    def test_art_delivery_uses_saved_clause_and_date_fallback(self):
        memo = {"contributorRole": "illustrator", "illustratorSketchDate": "2026-09-29",
                "illustratorFinalDate": "2026-12-31"}
        expected = ("The Illustrator shall create and furnish sketches to the Publisher by 2026-09-29. "
                    "The Illustrator shall create and furnish to the Publisher on or before 2026-12-31, "
                    "one (1) complete set of the Artwork in final form.")
        self.assertEqual(ns['_build_art_delivery_block'](memo), expected)
        self.assertEqual(ns['_build_manuscript_delivery_block'](memo), expected)
        memo['deliveryClause'] = expected + " Custom approved wording."
        self.assertEqual(ns['_build_art_delivery_block'](memo), memo['deliveryClause'])
        self.assertEqual(ns['_build_art_delivery_block']({'delivery_clause': expected}), expected)
        self.assertEqual(ns['_default_mapping']('illustrator')['ART_DELIVERY_BLOCK'], '__ART_DELIVERY_BLOCK__')
        self.assertIn('upon signing', ns['_build_manuscript_delivery_block']({'contributorRole': 'author'}))

    def test_agency_full_and_separate_address(self):
        memo = {'agency_street': '11 Briarwood Lane', 'agency_city': 'West Tisbury',
                'agency_state': 'MA', 'agency_zip': '02575', 'agency_country': 'USA'}
        mapping = ns['_default_mapping']('illustrator')
        self.assertEqual(ns['_get_value_from_memo'](memo, mapping['Agency_Address']),
                         '11 Briarwood Lane, West Tisbury, MA 02575, USA')
        self.assertEqual(ns['_get_value_from_memo'](memo, mapping['Agency_Street_Address']), '11 Briarwood Lane')
        self.assertEqual(ns['_get_value_from_memo']({'agency_city': 'Boston'}, 'agency_address'), 'Boston')
        self.assertEqual(ns['_get_value_from_memo']({}, 'agency_address'), '')

    def test_email_inside_hyperlink_is_processed_by_generator(self):
        # Exercise the actual generator's paragraph handler, including its guards.
        handler = next(n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef) and n.name == 'replace_in_paragraph')
        context = dict(ns, has_boardbook=True, SUBRIGHT_TOKENS=[])
        memo = {'author': {'email': 'author@example.test'}}
        context['values'] = {'Author_Email': ns['_get_value_from_memo'](memo, 'author_email')}
        exec(compile(ast.Module(body=[handler], type_ignores=[]), str(SOURCE), 'exec'), context)
        doc = Document()
        p = doc.add_paragraph()
        hyperlink = OxmlElement('w:hyperlink')
        for text in ('{{Author_', 'Email}}'):
            run = OxmlElement('w:r')
            node = OxmlElement('w:t')
            node.text = text
            run.append(node)
            hyperlink.append(run)
        p._p.append(hyperlink)
        self.assertFalse(p.runs)
        context['replace_in_paragraph'](p)
        self.assertEqual(''.join(p._p.xpath('.//w:t/text()')), 'author@example.test')

    def test_royalty_points_font_and_justification_round_trip(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from routers.contract_royalties import render_royalty_sections
        memo = {'royalties': {'illustrator': {'first_rights': [
            {'format': 'Hardcover', 'mode': 'flat', 'flat_rate_percent': 10, 'base': 'list_price'},
            {'format': 'Paperback', 'mode': 'flat', 'flat_rate_percent': 8, 'base': 'list_price'}]}}}
        for block in (True, False):
            doc = Document()
            heading = doc.add_paragraph('8. Royalties.')
            heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
            num = heading._p.get_or_add_pPr().get_or_add_numPr()
            num.get_or_add_ilvl().val = 0
            num.get_or_add_numId().val = 1
            slot = doc.add_paragraph('{{ROYALTIES_BLOCK}}' if block else '{{Hardcover_1}}')
            slot.alignment = WD_ALIGN_PARAGRAPH.LEFT
            slot.runs[0].font.name = 'Arial'
            slot.runs[0].font.size = Pt(9)
            render_royalty_sections(doc, memo, 'illustrator', insert_block=ns['_insert_numbered_contract_block'])
            output = io.BytesIO()
            doc.save(output)
            output.seek(0)
            result = Document(output)
            self.assertEqual(result.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)
            self.assertEqual(len(result.paragraphs[1:]), 2)
            for p in result.paragraphs[1:]:
                self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
                for run in p.runs:
                    if run.text:
                        self.assertEqual(run.font.name, 'Times New Roman')
                        self.assertEqual(run.font.size, Pt(12))
                        self.assertEqual(run._r.rPr.rFonts.get(qn('w:eastAsia')), 'Times New Roman')

    def test_headers_footers_and_nested_tables_round_trip(self):
        doc = Document()
        doc.add_paragraph('{{Illustrator_Name}}')
        section = doc.sections[0]
        for story in (section.header, section.first_page_header, section.even_page_header,
                      section.footer, section.first_page_footer, section.even_page_footer):
            story.paragraphs[0].text = '{{Illustrator_Email}}'
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.add_table(rows=1, cols=1).cell(0, 0).text = '{{Agency_Name}}'
        doc.add_section()  # linked stories must not duplicate or create extra definitions
        paragraphs = list(ns['_iter_contract_paragraphs'](doc))
        self.assertEqual(len(paragraphs), len({p._p for p in paragraphs}))
        values = {'Illustrator_Name': 'Iris Artist', 'Illustrator_Email': 'iris@example.test', 'Agency_Name': 'Art Agency'}
        for p in paragraphs:
            ns['_replace_inline_tokens'](p, values)
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        result = Document(output)
        text = '\n'.join(p.text for p in ns['_iter_contract_paragraphs'](result))
        self.assertNotIn('{{', text)
        self.assertEqual(text.count('iris@example.test'), 6)
        self.assertIn('Iris Artist', text)
        self.assertIn('Art Agency', text)
        self.assertTrue(result.sections[1].header.is_linked_to_previous)

if __name__ == '__main__':
    unittest.main()
